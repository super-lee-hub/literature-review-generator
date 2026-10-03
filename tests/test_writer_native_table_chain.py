from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from docx_writer import rebuild_review_docx_from_structured_artifacts, scan_docx_for_unresolved_citation_tokens
from services.citation_manifest import build_citation_manifest_from_review_draft
from services.job_workspace import JobWorkspace
from services.review_draft import build_review_draft, iter_review_text_blocks
from services.review_generation_service import ReviewGenerationService
from services.settings import ApplicationSettings
from runtime.reconcile import DEFAULT_SCHEMA_VALIDATORS
from runtime.artifact_validators import _validate_review_json
from tests.test_writer_source_inventory import (
    JOB_ID, PLANNED_CLAIM, QUALIFIER_CLAIM_ID, QUALIFIER_TEXT, SOURCE_SUMMARY, _registry_with_inventory,
)


def _plan():
    return {
        "schema_version": "writer_table_plan/v1",
        "tables": [{
            "table_id": "results", "headers": ["Evidence", "Finding"],
            "rows": [
                {"row_id": "finding", "cells": [
                    {"static_text": "Result"}, {"planned_claim_index": 0, "source_claim_id": None},
                ]},
                {"row_id": "condition", "cells": [
                    {"static_text": "Condition"}, {"planned_claim_index": 0, "source_claim_id": QUALIFIER_CLAIM_ID},
                ]},
            ],
        }],
    }


def _service(tmp_path, calls):
    workspace = JobWorkspace.create(str(tmp_path / "output"), "native-table", JOB_ID)
    registry, _inventory, packet, _layers, _views = _registry_with_inventory(Path(workspace.root_dir))

    def writer(**kwargs):
        scope = json.loads(kwargs["prompt_text"])["writer_task_scope"]
        calls.append(scope)
        return {"status": "success", "usage_status": "provider_not_supported", "content": {
            "blocks": [{
                "writer_task_id": task["writer_task_id"],
                "writer_output_unit_id": unit["writer_output_unit_id"],
                "writer_task_basis_hash": scope["writer_task_basis_hash"],
                "text": f"{PLANNED_CLAIM if unit['unit_kind'] == 'planned_claim' else QUALIFIER_TEXT} [[cite_ref:R001]]",
            } for task in scope["tasks"] for unit in task["output_units"] if unit["required"]],
            "task_dispositions": [{
                "writer_task_id": task["writer_task_id"], "writer_task_basis_hash": scope["writer_task_basis_hash"],
                "disposition": "covered",
            } for task in scope["tasks"]],
        }}

    service = ReviewGenerationService(
        job_id=JOB_ID, attempt_id="native-table-attempt", workspace=workspace, artifact_registry=registry,
        settings=ApplicationSettings.from_config({"Writer_API": {
            "api_key": "synthetic", "model": "writer", "api_base": "https://writer.test/v1",
        }}), summaries=[deepcopy(SOURCE_SUMMARY)], writer=writer,
    )
    outline = {"title": "Bounded table review", "sections": [{
        "section_id": packet["section_id"], "title": "Results", "writer_table_plan": _plan(),
    }]}
    return service, packet, outline


def _manifest(draft, result, output_path):
    return build_citation_manifest_from_review_draft(
        job_id=JOB_ID, project_name="native-table", manifest_id="native-table-manifest",
        review_draft_path="local:draft", review_word_path=str(output_path), review_draft=draft,
        paper_summaries=[deepcopy(SOURCE_SUMMARY)], citation_ref_catalog=result.citation_ref_catalog,
        citation_ref_catalog_path=result.citation_ref_catalog_path,
    )


def test_native_table_keeps_source_units_through_writer_draft_manifest_and_docx(tmp_path):
    calls = []
    service, packet, outline = _service(tmp_path, calls)
    result = service.run(outline_payload=outline, evidence_packets=[packet])
    assert len(calls) == 1
    section = result.sections[0]
    assert section["writer_task_scope"]["schema_version"] == "writer_task_scope/v2"
    assert section["writer_task_scope"]["min_required_output_units"] == 2
    table = section["blocks"][0]
    assert table["block_kind"] == "table" and table["text"] == ""
    cells = [row["cells"][1] for row in table["rows"]]
    assert len({cell["writer_output_unit_id"] for cell in cells}) == 2
    assert all(QUALIFIER_CLAIM_ID in cell["required_source_context"]["source_claim_ids"] for cell in cells)
    assert {block["block_id"] for block in iter_review_text_blocks(section)} == {cell["block_id"] for cell in cells}

    output_path = tmp_path / "native-table.docx"
    draft = build_review_draft(
        job_id=JOB_ID, project_name="native-table", draft_id="native-table-draft", title="Bounded table review",
        outline_artifact_id="outline-v3:final_outline", outline_source_path="local:outline",
        summary_file="local:summary", review_word_path=str(output_path), sections=result.sections,
        references=[], generation_mode="review_v3", citation_ref_catalog=result.citation_ref_catalog,
    ).to_dict()
    draft_path = tmp_path / "native-table-draft.json"
    draft_path.write_text(json.dumps(draft, ensure_ascii=False), encoding="utf-8")
    draft_record = service.registry.register_file(
        artifact_id="native-table-draft", artifact_role="review_draft", artifact_type="review_draft",
        artifact_version="v3", producer="test", path=draft_path,
    )
    DEFAULT_SCHEMA_VALIDATORS["review_draft"](draft_record, draft_path)
    _validate_review_json(draft_record, draft_path, draft)
    manifest = _manifest(draft, result, output_path).to_dict()
    claim_units = [unit for bundle in manifest["citation_sets"] for unit in bundle["claim_units"]]
    assert len(claim_units) == len(manifest["occurrences"]) == 2
    assert {unit["block_id"] for unit in claim_units} == {cell["block_id"] for cell in cells}
    assert {unit["writer_output_unit_id"] for unit in claim_units} == {cell["writer_output_unit_id"] for cell in cells}
    assert all(unit["table_id"] == "results" and unit["cell_id"] for unit in claim_units)
    assert all(QUALIFIER_CLAIM_ID in unit["required_source_context"]["source_claim_ids"] for unit in claim_units)
    rebuild_review_docx_from_structured_artifacts(SimpleNamespace(logger=None), draft, manifest, str(output_path))
    document = Document(str(output_path))
    assert len(document.tables) == 1 and len(document.tables[0].rows) == 3
    assert document.tables[0].rows[0].cells[0].text == "Evidence"
    assert QUALIFIER_TEXT in document.tables[0].rows[2].cells[1].text
    assert scan_docx_for_unresolved_citation_tokens(str(output_path), manifest)["passed"] is True

    for mutation in ("header", "static_fact", "source_context", "drop_cell", "duplicate_unit", "parent_text", "foreign_metadata", "stale_span"):
        tampered = deepcopy(draft)
        target = tampered["content"]["sections"][0]["blocks"][0]
        if mutation == "header":
            target["headers"][0] = "Different evidence"
        elif mutation == "static_fact":
            target["rows"][0]["cells"][0]["text"] = "The treatment works everywhere"
        elif mutation == "source_context":
            context = target["rows"][0]["cells"][1]["required_source_context"]
            context["source_claim_ids"] = [identifier for identifier in context["source_claim_ids"] if identifier != QUALIFIER_CLAIM_ID]
        elif mutation == "drop_cell":
            target["rows"][1]["cells"].pop()
        elif mutation == "duplicate_unit":
            target["rows"][1]["cells"][1] = deepcopy(target["rows"][0]["cells"][1])
        elif mutation == "parent_text":
            target["text"] = "A hidden unbound conclusion [[cite_ref:R001]]."
        elif mutation == "foreign_metadata":
            citation = target["rows"][0]["cells"][1]["citations"][0]
            citation.update({"ref_id": "R999", "citation_token": "[[cite_ref:R999]]", "raw_text": "[[cite_ref:R999]]"})
        else:
            citation = target["rows"][0]["cells"][1]["citations"][0]
            citation.update({"span_start": 0, "span_end": 16})
        try:
            _manifest(tampered, result, output_path)
        except ValueError:
            pass
        else:
            pytest.fail(f"Citation manifest accepted table mutation: {mutation}")
        with pytest.raises(ValueError):
            _validate_review_json(draft_record, draft_path, tampered)


def test_table_plan_rejects_foreign_source_before_writer_dispatch(tmp_path):
    calls = []
    service, packet, outline = _service(tmp_path, calls)
    outline["sections"][0]["writer_table_plan"]["tables"][0]["rows"][1]["cells"][1]["source_claim_id"] = "invented:claim"
    with pytest.raises(ValueError):
        service.run(outline_payload=outline, evidence_packets=[packet])
    assert calls == []
    assert service.registry.get(f"review-section:{packet['section_id']}") is None
