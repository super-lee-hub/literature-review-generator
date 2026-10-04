from __future__ import annotations

import copy
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Mapping

import pytest
from docx import Document
from docx.oxml.ns import qn

from docx_writer import (
    append_review_section_blocks_to_word_document,
    rebuild_review_docx_from_structured_artifacts,
    scan_docx_for_unresolved_citation_tokens,
)
from services.artifact_registry import ArtifactRegistry
from services.citation_manifest import build_citation_manifest_from_review_draft
from services.job_workspace import JobWorkspace
from services.review_draft import build_review_draft
from services.review_generation_service import ReviewGenerationService
from services.settings import ApplicationSettings
from tests import test_current_review_generation as review_generation_helpers
from tests.writer_source_fixture import bind_production_writer_sources
from tests.test_docx_citation_renderer import _manifest


def test_review_service_projects_source_bound_cjk_sentence_into_multipage_native_table(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    long_finding = ("用于分页测试的中文内容，" * 250) + "本句仅用于检验原生表格分页。"
    qualifier = "本结果仅为合成排版测试文本，不能外推为真实研究结论。"
    assert 3000 <= len(long_finding) < 4096
    assert long_finding.count("。") == 1

    canonical_summary = copy.deepcopy(review_generation_helpers._canonical_summary())
    canonical_summary["core_analysis"]["findings"] = long_finding
    canonical_summary["core_analysis"]["key_points"] = [long_finding]
    canonical_summary["core_analysis"]["limitations"] = qualifier
    monkeypatch.setattr(
        review_generation_helpers,
        "_canonical_summary",
        lambda: copy.deepcopy(canonical_summary),
    )

    summary, _source_pdf = review_generation_helpers._stage1_summary(tmp_path)
    paper_key = str(summary["paper_info"]["canonical_paper_key"])
    packet: dict[str, Any] = {
        "section_id": "section:results",
        "section_goal": "展示来源约束下的长中文结果",
        "planned_claims": [long_finding],
        "paper_keys": [paper_key],
        "source_summary_hashes": [],
        "retrieval_provenance": {
            "source": "stage1_summary",
            "paper_keys": [paper_key],
        },
    }

    seen_writer_scope: list[Mapping[str, Any]] = []

    def writer(**kwargs: Any) -> Mapping[str, Any]:
        prompt = json.loads(str(kwargs["prompt_text"]))
        scope = prompt["writer_task_scope"]
        seen_writer_scope.append(scope)
        assert scope["schema_version"] == "writer_task_scope_wire/v2"
        task = scope["tasks"][0]
        primary_unit = next(
            unit for unit in task["output_units"]
            if unit["unit_kind"] == "planned_claim"
        )
        qualifier_claim = next(
            scope["evidence_store"][reference]["value"]
            for row in task["canonical_source_bundle_refs"]
            for reference in row["source_claims_refs"]
            if scope["evidence_store"][reference]["value"].get("text") == qualifier
        )
        assert qualifier_claim["claim_id"] in task["qualifier_source_claim_ids"]
        qualifier_unit = next(
            unit for unit in task["output_units"]
            if unit["unit_kind"] == "source_claim_expansion"
            and unit["source_claim_id"] == qualifier_claim["claim_id"]
        )
        assert primary_unit["required"] is True
        assert qualifier_unit["required"] is False
        citation = f"[[cite_ref:{task['allowed_ref_ids'][0]}]]"
        primary_text = f"{long_finding[:-1]} {citation}{long_finding[-1]}"
        qualifier_text = f"{qualifier[:-1]} {citation}{qualifier[-1]}"
        basis = str(scope["writer_task_basis_hash"])
        return {
            "status": "success",
            "content": {
                "blocks": [
                    {
                        "writer_task_id": task["writer_task_id"],
                        "writer_output_unit_id": primary_unit["writer_output_unit_id"],
                        "writer_task_basis_hash": basis,
                        "text": primary_text,
                    },
                    {
                        "writer_task_id": task["writer_task_id"],
                        "writer_output_unit_id": qualifier_unit["writer_output_unit_id"],
                        "writer_task_basis_hash": basis,
                        "text": qualifier_text,
                    },
                ],
                "task_dispositions": [{
                    "writer_task_id": task["writer_task_id"],
                    "writer_task_basis_hash": basis,
                    "disposition": "covered",
                }],
            },
        }

    workspace = JobWorkspace.create(
        str(tmp_path / "review-output"),
        "current",
        job_id="docx-table-job",
    )
    registry = ArtifactRegistry(workspace.paths.registry_path, workspace.job_id)
    service = ReviewGenerationService(
        job_id=workspace.job_id,
        attempt_id="docx-table-attempt",
        workspace=workspace,
        artifact_registry=registry,
        settings=ApplicationSettings.from_config({
            "Writer_API": {
                "api_key": "writer-test",
                "model": "writer-test",
                "api_base": "https://writer.test/v1",
                "max_context_tokens": "128000",
                "max_output_tokens": "4096",
            },
        }),
        summaries=[summary],
        writer=writer,
    )
    bind_production_writer_sources(service, [packet])
    result = service.run(
        outline_payload={
            "title": "本地排版验证",
            "sections": [{
                "section_id": "section:results",
                "title": "中文结果",
                "goal": "展示来源约束下的长中文结果",
                "writer_table_plan": {
                    "schema_version": "writer_table_plan/v1",
                    "tables": [{
                        "table_id": "results_table",
                        "headers": ["指标", "结果"],
                        "rows": [{
                            "row_id": "primary_finding",
                            "cells": [
                                {"static_text": "研究发现"},
                                {"planned_claim_index": 0, "source_claim_id": None},
                            ],
                        }],
                    }],
                },
            }],
        },
        evidence_packets=[packet],
    )

    assert len(seen_writer_scope) == 1
    assert seen_writer_scope[0]["source_inventory_binding"]["artifact_id"]
    task_source_hashes = {
        row["source_summary_hash"]
        for row in result.sections[0]["writer_task_scope"]["tasks"][0]["source_evidence"]
    }
    assert packet["source_summary_hashes"] == sorted(task_source_hashes)
    assert all(len(value) == 64 for value in packet["source_summary_hashes"])
    assert len(result.sections) == 1
    section = result.sections[0]
    table_block = next(block for block in section["blocks"] if block["block_kind"] == "table")
    assert table_block.get("text", "") == ""
    assert table_block["headers"] == ["指标", "结果"]
    assert len(table_block["rows"]) == 1
    label_cell, factual_cell = table_block["rows"][0]["cells"]
    assert label_cell["text"] == "研究发现"
    assert label_cell["cell_kind"] == "static_label"
    assert factual_cell["cell_kind"] == "factual_output_unit"
    assert factual_cell["text"] == f"{long_finding[:-1]} [[cite_ref:R001]]{long_finding[-1]}"
    assert factual_cell["block_id"]
    assert factual_cell["writer_task_id"]
    assert factual_cell["writer_output_unit_id"]
    assert factual_cell["writer_task_basis_hash"] == table_block["writer_task_basis_hash"]
    assert factual_cell["required_source_context"]["source_claim_ids"]
    assert factual_cell["required_source_context"]["qualifier_source_claim_ids"]
    qualifier_block = next(
        block for block in section["blocks"]
        if block["block_kind"] == "paragraph"
        and block["text"].startswith(qualifier[:-1])
        and block["text"].endswith(qualifier[-1])
    )
    assert qualifier_block["block_id"]
    assert qualifier_block["writer_output_unit_id"]
    assert qualifier_block["required_source_context"]["qualifier_source_claim_ids"]

    output_path = tmp_path / "review-with-native-table.docx"
    review_draft = build_review_draft(
        job_id=workspace.job_id,
        project_name=workspace.project_name,
        draft_id="docx-table-draft",
        outline_artifact_id="outline-v3:final_outline",
        title="本地排版验证",
        outline_source_path="local:docx-table-outline",
        summary_file="local:docx-table-summary",
        review_word_path=str(output_path),
        sections=result.sections,
        references=[],
        generation_mode="review_v3",
        paper_summaries=[summary],
        citation_ref_catalog=result.citation_ref_catalog,
        citation_ref_catalog_path=result.citation_ref_catalog_path,
        citation_ref_catalog_hash=str(
            result.citation_ref_catalog.get("catalog_hash") or ""
        ),
    )
    review_draft_payload = review_draft.to_dict()
    draft_section = review_draft_payload["content"]["sections"][0]
    draft_table = next(
        block for block in draft_section["blocks"]
        if block.get("table_layout_schema_version") == "writer_table_layout/v1"
    )
    assert draft_table["text"] == ""
    draft_fact_cell = draft_table["rows"][0]["cells"][1]
    assert draft_fact_cell["block_id"] == factual_cell["block_id"]
    assert draft_fact_cell["citations"][0]["ref_id"] == "R001"
    citation_manifest = build_citation_manifest_from_review_draft(
        job_id=workspace.job_id,
        project_name=workspace.project_name,
        manifest_id="docx-table-citation-manifest",
        review_draft_path=str(tmp_path / "review-draft.json"),
        review_word_path=str(output_path),
        review_draft=review_draft_payload,
        paper_summaries=[summary],
        citation_ref_catalog=result.citation_ref_catalog,
        citation_ref_catalog_path=result.citation_ref_catalog_path,
        citation_ref_catalog_hash=str(
            result.citation_ref_catalog.get("catalog_hash") or ""
        ),
    )
    manifest_payload = citation_manifest.to_dict()
    assert any(
        occurrence["block_id"] == factual_cell["block_id"]
        and occurrence["ref_id"] == "R001"
        for occurrence in manifest_payload["occurrences"]
    )
    assert any(
        occurrence["block_id"] == qualifier_block["block_id"]
        and occurrence["ref_id"] == "R001"
        for occurrence in manifest_payload["occurrences"]
    )

    rebuild_review_docx_from_structured_artifacts(
        SimpleNamespace(logger=None),
        review_draft_payload,
        manifest_payload,
        str(output_path),
    )

    document = Document(str(output_path))
    assert document.sections[0].header.paragraphs[0].text == "本地排版验证"
    assert document.sections[0].footer._element.xml.count("PAGE") == 1
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.rows[0]._tr.trPr is not None
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    table_text = "\n".join(
        cell.text for row in table.rows for cell in row.cells
    )
    assert "指标" in table_text
    assert "研究发现" in table_text
    assert len(table.cell(1, 1).text) > 2500
    assert table.cell(1, 1).text.startswith(long_finding[:200])
    assert long_finding[1200:1400] in table.cell(1, 1).text
    assert table.cell(1, 1).text.endswith("。")
    assert "[[cite_ref:" not in table_text
    assert "(" in table_text and ")" in table_text
    assert any(
        paragraph.text.startswith(qualifier[:-1])
        and paragraph.text.endswith(qualifier[-1])
        for paragraph in document.paragraphs
    )

    scan = scan_docx_for_unresolved_citation_tokens(
        str(output_path),
        manifest_payload,
    )
    assert scan["passed"] is True, scan


def test_legacy_review_draft_uses_default_docx_header_without_duplicate_footer(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "legacy-review.docx"
    review_draft = {
        "draft_identity": {"draft_id": "legacy-draft"},
        "content": {
            "sections": [{
                "section_number": 1,
                "section_title": "Results",
                "blocks": [{"block_kind": "paragraph", "text": "A local result."}],
            }],
            "references": [],
        },
    }

    rebuild_review_docx_from_structured_artifacts(
        SimpleNamespace(logger=None),
        review_draft,
        {"bibliography": []},
        str(output_path),
    )

    document = Document(str(output_path))
    assert document.sections[0].header.paragraphs[0].text == "Literature Review"
    assert document.sections[0].footer._element.xml.count("PAGE") == 1


def test_explicit_table_block_is_native_and_malformed_table_fails_closed(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "explicit-table.docx"
    appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [{
            "block_id": "s1_b1",
            "block_kind": "table",
            "text": "| Measure | Result |\n| --- | --- |\n| Sample | 24 |",
        }],
        str(output_path),
        citation_manifest={},
    )

    assert appended is True
    assert len(Document(str(output_path)).tables) == 1

    malformed_path = tmp_path / "malformed-table.docx"
    malformed_appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [{
            "block_id": "s1_b2",
            "block_kind": "table",
            "text": "| Measure | Result |\nThis is missing a separator row.",
        }],
        str(malformed_path),
        citation_manifest={},
    )

    assert malformed_appended is False
    assert not malformed_path.exists()

    empty_path = tmp_path / "empty-table.docx"
    empty_appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [{"block_id": "s1_b3", "block_kind": "table", "text": "  "}],
        str(empty_path),
        citation_manifest={},
    )
    assert empty_appended is False
    assert not empty_path.exists()


def test_ordinary_prose_with_pipe_remains_a_paragraph(tmp_path: Path) -> None:
    output_path = tmp_path / "ordinary-pipe-prose.docx"
    prose = "The expression A | B is a comparison, not a table."

    appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [{
            "block_id": "s1_b1",
            "block_kind": "paragraph",
            "text": prose,
        }],
        str(output_path),
        citation_manifest={},
    )

    assert appended is True
    document = Document(str(output_path))
    assert document.tables == []
    assert any(paragraph.text == prose for paragraph in document.paragraphs)


def test_same_reference_keeps_block_and_cell_specific_modes_and_locators(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "table-citation-modes.docx"
    manifest = _manifest()
    manifest["occurrences"] = [
        {"ref_id": "R006", "paper_id": "10.1/a", "section_number": 1,
         "block_id": "s1_b1", "mode": "narrative", "locator": "p. 1"},
        {"ref_id": "R006", "paper_id": "10.1/a", "section_number": 1,
         "block_id": "s1_b2", "mode": "parenthetical", "locator": "p. 2"},
        {"ref_id": "R006", "paper_id": "10.1/a", "section_number": 1,
         "block_id": "s1_b2", "mode": "narrative", "locator": "p. 3"},
    ]
    appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [
            {"block_id": "s1_b1", "block_kind": "paragraph",
             "text": "Opening claim [[cite_ref:R006]]."},
            {"block_id": "s1_b2", "block_kind": "paragraph",
             "text": "| Item | Finding |\n| --- | --- |\n"
                     "| First | [[cite_ref:R006]] |\n"
                     "| Second | [[cite_ref:R006]] |"},
        ],
        str(output_path),
        citation_manifest=manifest,
    )

    assert appended is True
    document = Document(str(output_path))
    assert any("Chen (2024, p. 1)" in paragraph.text for paragraph in document.paragraphs)
    assert "(Chen, 2024, p. 2)" in document.tables[0].cell(1, 1).text
    assert "Chen (2024, p. 3)" in document.tables[0].cell(2, 1).text
