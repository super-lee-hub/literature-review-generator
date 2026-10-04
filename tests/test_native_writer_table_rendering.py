from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from docx_writer import (
    append_review_section_blocks_to_word_document,
    rebuild_review_docx_from_structured_artifacts,
)


_BASIS_HASH = "a" * 64


def _native_table_block() -> dict[str, object]:
    source_context = {
        "source_claim_ids": ["claim:1"],
        "evidence_ids": ["evidence:1"],
        "source_field_ids": ["field:sample"],
        "qualifier_source_claim_ids": [],
        "qualifier_evidence_ids": [],
        "qualifier_source_field_ids": [],
    }
    return {
        "block_id": "writer_table_study-results",
        "block_kind": "table",
        "table_layout_schema_version": "writer_table_layout/v1",
        "table_id": "study-results",
        "writer_task_basis_hash": _BASIS_HASH,
        "headers": ["Measure", "Finding"],
        "rows": [
            {
                "row_id": "sample",
                "cells": [
                    {
                        "cell_id": "study-results:sample:0",
                        "block_id": "writer_static_cell_sample",
                        "cell_kind": "static_label",
                        "text": "Sample",
                        "source_validation_status": "caller_allowlisted_nonfactual_label",
                    },
                    {
                        "cell_id": "study-results:sample:1",
                        "block_id": "writer_cell_unit0001",
                        "cell_kind": "factual_output_unit",
                        "writer_task_id": "task0001",
                        "writer_output_unit_id": "unit0001",
                        "writer_task_basis_hash": _BASIS_HASH,
                        "text": "The sample had 24 participants [[cite_ref:R001]].",
                        "allowed_ref_ids": ["R001"],
                        "required_source_context": source_context,
                        "source_validation_status": "canonical_source_inventory_verified",
                    },
                ],
            },
            {
                "row_id": "setting",
                "cells": [
                    {
                        "cell_id": "study-results:setting:0",
                        "block_id": "writer_static_cell_setting",
                        "cell_kind": "static_label",
                        "text": "Setting",
                        "source_validation_status": "caller_allowlisted_nonfactual_label",
                    },
                    {
                        "cell_id": "study-results:setting:1",
                        "block_id": "writer_cell_unit0002",
                        "cell_kind": "factual_output_unit",
                        "writer_task_id": "task0001",
                        "writer_output_unit_id": "unit0002",
                        "writer_task_basis_hash": _BASIS_HASH,
                        "text": "The study was conducted online [[cite_ref:R001]].",
                        "allowed_ref_ids": ["R001"],
                        "required_source_context": source_context,
                        "source_validation_status": "canonical_source_inventory_verified",
                    },
                ],
            },
        ],
    }


def _citation_manifest() -> dict[str, object]:
    return {
        "paper_entries": [
            {
                "paper_id": "10.1/a",
                "paper_key": "10.1/a",
                "title": "A Study",
                "authors": ["Lanfei Chen"],
                "year": "2024",
            },
        ],
        "occurrences": [
            {
                "ref_id": "R001",
                "paper_id": "10.1/a",
                "section_number": 1,
                "block_id": "writer_cell_unit0001",
                "mode": "narrative",
                "locator": "p. 7",
            },
            {
                "ref_id": "R001",
                "paper_id": "10.1/a",
                "section_number": 1,
                "block_id": "writer_cell_unit0002",
                "mode": "parenthetical",
                "locator": "p. 9",
            },
        ],
        "bibliography": [],
    }


def test_native_writer_table_rebuilds_with_cell_citations_and_repeating_header(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "native-writer-table.docx"
    review_draft = {
        "draft_identity": {"title": "Native table rendering"},
        "content": {
            "sections": [
                {
                    "section_number": 1,
                    "section_title": "Results",
                    "blocks": [_native_table_block()],
                },
            ],
        },
    }

    rebuild_review_docx_from_structured_artifacts(
        SimpleNamespace(logger=None),
        review_draft,
        _citation_manifest(),
        str(output_path),
    )

    document = Document(str(output_path))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert [cell.text for cell in table.rows[0].cells] == ["Measure", "Finding"]
    assert [cell.text for cell in table.rows[1].cells] == [
        "Sample",
        "The sample had 24 participants Chen (2024, p. 7).",
    ]
    assert [cell.text for cell in table.rows[2].cells] == [
        "Setting",
        "The study was conducted online (Chen, 2024, p. 9).",
    ]
    assert "tblHeader" in table.rows[0]._tr.trPr.xml
    assert "[[cite_ref:" not in "\n".join(
        cell.text for row in table.rows for cell in row.cells
    )


def test_source_bound_paragraph_and_native_table_render_in_the_same_section(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "mixed-source-bound-review.docx"
    manifest = _citation_manifest()
    occurrences = manifest["occurrences"]
    assert isinstance(occurrences, list)
    manifest["occurrences"] = [
        {
            "ref_id": "R001",
            "paper_id": "10.1/a",
            "section_number": 1,
            "block_id": "writer_unit_paragraph0001",
            "mode": "parenthetical",
        },
        *occurrences,
    ]
    blocks = [
        {
            "block_id": "writer_unit_paragraph0001",
            "block_kind": "paragraph",
            "writer_task_id": "task0001",
            "writer_output_unit_id": "paragraph0001",
            "writer_task_basis_hash": _BASIS_HASH,
            "text": "Introductory result [[cite_ref:R001]].",
        },
        _native_table_block(),
    ]

    appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        blocks,
        str(output_path),
        citation_manifest=manifest,
    )

    assert appended is True
    document = Document(str(output_path))
    assert len(document.tables) == 1
    assert any("Introductory result (Chen, 2024)." == p.text for p in document.paragraphs)
    assert document.tables[0].cell(1, 1).text.endswith("Chen (2024, p. 7).")


@pytest.mark.parametrize(
    "mutate",
    [
        lambda block: block.pop("writer_task_basis_hash"),
        lambda block: block["rows"][0]["cells"].pop(),
        lambda block: block["rows"][0]["cells"][0].update(
            text="Outcome [[cite_ref:R001]]"
        ),
        lambda block: block["rows"][0]["cells"][1].pop("required_source_context"),
        lambda block: block.update(text="unexpected parent text"),
        lambda block: block.update(table_layout_schema_version="writer_table_layout/v2"),
        lambda block: block["rows"][0]["cells"][1].update(
            allowed_ref_ids=["R002"]
        ),
    ],
    ids=[
        "missing-basis-hash",
        "row-dimension-mismatch",
        "static-label-citation",
        "missing-source-context",
        "parent-text",
        "unsupported-schema",
        "foreign-citation",
    ],
)
def test_malformed_native_writer_table_fails_closed_before_writing(
    tmp_path: Path,
    mutate: object,
) -> None:
    block = _native_table_block()
    mutate(block)  # type: ignore[operator]
    output_path = tmp_path / "malformed-native-table.docx"

    appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [block],
        str(output_path),
        citation_manifest=_citation_manifest(),
    )

    assert appended is False
    assert not output_path.exists()


def test_rebuild_preserves_existing_docx_when_native_table_shape_is_invalid(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "existing-review.docx"
    existing = Document()
    existing.add_paragraph("Existing complete document")
    existing.save(str(output_path))
    bad_block = _native_table_block()
    bad_block["rows"][0]["cells"].pop()  # type: ignore[index]

    with pytest.raises(ValueError, match="dimensions"):
        rebuild_review_docx_from_structured_artifacts(
            SimpleNamespace(logger=None),
            {
                "content": {
                    "sections": [
                        {
                            "section_number": 1,
                            "section_title": "Results",
                            "blocks": [bad_block],
                        },
                    ],
                },
            },
            _citation_manifest(),
            str(output_path),
        )

    assert Document(str(output_path)).paragraphs[0].text == "Existing complete document"


def test_legacy_markdown_table_keeps_parent_block_citation_behavior(
    tmp_path: Path,
) -> None:
    output_path = tmp_path / "legacy-markdown-table.docx"
    manifest = _citation_manifest()
    manifest["occurrences"] = [
        {
            "ref_id": "R001",
            "paper_id": "10.1/a",
            "section_number": 1,
            "block_id": "legacy_table_block",
            "mode": "narrative",
            "locator": "p. 3",
        },
    ]
    markdown_table = "\n".join(
        (
            "| Measure | Finding |",
            "| --- | --- |",
            "| Sample | 24 participants [[cite_ref:R001]] |",
        )
    )

    appended = append_review_section_blocks_to_word_document(
        SimpleNamespace(logger=None),
        1,
        "Results",
        [
            {
                "block_id": "legacy_table_block",
                "block_kind": "paragraph",
                "text": markdown_table,
            },
        ],
        str(output_path),
        citation_manifest=manifest,
    )

    assert appended is True
    table = Document(str(output_path)).tables[0]
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.cell(1, 1).text == "24 participants Chen (2024, p. 3)"
