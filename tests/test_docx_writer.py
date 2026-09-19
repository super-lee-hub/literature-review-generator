"""Tests for the current structured-citation DOCX contract."""

from __future__ import annotations

import pytest

from docx_writer import (
    append_section_to_word_document,
    generate_apa_references_from_manifest,
    rebuild_final_docx_from_manifest,
    rebuild_review_docx_from_structured_artifacts,
    scan_docx_for_unresolved_citation_tokens,
)


class MockGenerator:
    def __init__(self) -> None:
        self.logger = type("Logger", (), {"error": lambda _self, _message: None})()


def _manifest() -> dict[str, object]:
    return {
        "paper_entries": [
            {
                "paper_id": "paper_1",
                "paper_key": "paper_1",
                "title": "Structured Paper One",
                "authors": ["Alice Smith"],
                "year": "2024",
            },
            {
                "paper_id": "paper_2",
                "paper_key": "paper_2",
                "title": "Structured Paper Two",
                "authors": ["Bob Jones"],
                "year": "2025",
            },
        ],
        "occurrences": [
            {"ref_id": "R001", "paper_id": "paper_1", "paper_key": "paper_1"},
            {"ref_id": "R002", "paper_id": "paper_2", "paper_key": "paper_2"},
        ],
        "bibliography": [
            {
                "entry_id": "bib_001",
                "paper_id": "paper_1",
                "paper_key": "paper_1",
                "citation_text": "Alice Smith (2024). Structured Paper One.",
                "is_cited": True,
            },
            {
                "entry_id": "bib_002",
                "paper_id": "paper_2",
                "paper_key": "paper_2",
                "citation_text": "Bob Jones (2025). Structured Paper Two.",
                "is_cited": True,
            },
        ],
    }


def test_manifest_first_bibliography() -> None:
    references = generate_apa_references_from_manifest(
        {
            "bibliography": [
                {
                    "citation_text": "Author A, B. (2023). Test Paper 1.",
                    "is_cited": True,
                }
            ]
        },
        MockGenerator(),
    )

    assert references == ["Author A, B. (2023). Test Paper 1."]


def test_manifest_bibliography_contains_only_cited_entries() -> None:
    manifest = _manifest()
    manifest["bibliography"] = [
        manifest["bibliography"][0],
        {**manifest["bibliography"][1], "is_cited": False},
    ]

    references = generate_apa_references_from_manifest(manifest, MockGenerator())

    assert references == ["Alice Smith (2024). Structured Paper One."]


def test_rebuild_review_docx_requires_current_structured_citations(tmp_path) -> None:
    draft = {
        "content": {
            "sections": [
                {
                    "section_number": 1,
                    "section_title": "Intro",
                    "blocks": [{"text": "Structured claim [[cite_ref:R001]]."}],
                }
            ]
        }
    }
    output = tmp_path / "review.docx"
    rebuild_review_docx_from_structured_artifacts(
        MockGenerator(), draft, _manifest(), str(output)
    )

    from docx import Document

    text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
    assert "(Smith, 2024)" in text
    assert "Alice Smith (2024). Structured Paper One." in text

    draft["content"]["sections"][0]["blocks"][0]["text"] = (
        "Legacy mention [[cite:paper_1]]."
    )
    with pytest.raises(ValueError, match="section DOCX rendering failed"):
        rebuild_review_docx_from_structured_artifacts(
            MockGenerator(), draft, _manifest(), str(tmp_path / "legacy.docx")
        )


def test_final_docx_rebuild_scan_has_no_unresolved_tokens(tmp_path) -> None:
    draft = {
        "content": {
            "sections": [
                {
                    "section_number": 1,
                    "section_title": "Intro",
                    "blocks": [
                        {
                            "text": (
                                "Claim one [[cite_ref:R001]]. "
                                "Claim two [[cite_ref:R001, R002]]."
                            )
                        }
                    ],
                }
            ]
        }
    }
    output = tmp_path / "final.docx"
    scan_path = tmp_path / "scan.json"

    scan = rebuild_final_docx_from_manifest(
        MockGenerator(), draft, _manifest(), str(output), scan_report_path=str(scan_path)
    )

    assert scan["passed"] is True
    assert scan_path.exists()
    assert scan_docx_for_unresolved_citation_tokens(str(output), _manifest())[
        "unresolved_tokens"
    ] == []


def test_rebuild_review_docx_raises_when_section_append_fails(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("docx_writer.append_section_to_word_document", lambda *_args, **_kwargs: False)

    with pytest.raises(ValueError, match="section DOCX rendering failed"):
        rebuild_review_docx_from_structured_artifacts(
            MockGenerator(),
            {
                "content": {
                    "sections": [
                        {
                            "section_number": 1,
                            "section_title": "Intro",
                            "blocks": [{"text": "Claim [[cite_ref:R001]]."}],
                        }
                    ]
                }
            },
            _manifest(),
            str(tmp_path / "review.docx"),
        )


def test_docx_references_use_real_italics_and_hanging_indent(tmp_path) -> None:
    from docx import Document
    from zipfile import ZipFile

    manifest = _manifest()
    manifest["paper_entries"][0] = {
        **manifest["paper_entries"][0],
        "journal": "Journal of Testing",
        "volume": "12",
        "issue": "3",
        "pages": "44-59",
        "doi": "https://doi.org/10.1234/ABC",
    }
    manifest["bibliography"][0] = {
        **manifest["bibliography"][0],
        "citation_text": "Smith, A. (2024). Structured Paper One. Journal of Testing, 12(3), 44-59.",
    }
    output = tmp_path / "rich.docx"
    rebuild_review_docx_from_structured_artifacts(
        MockGenerator(),
        {
            "content": {
                "sections": [
                    {
                        "section_number": 1,
                        "section_title": "Intro",
                        "blocks": [{"text": "Claim [[cite_ref:R001]]."}],
                    }
                ]
            }
        },
        manifest,
        str(output),
    )

    doc = Document(str(output))
    reference = next(p for p in doc.paragraphs if p.text.startswith("Smith, A."))
    assert "*" not in reference.text
    assert reference.paragraph_format.first_line_indent is not None
    assert reference.paragraph_format.first_line_indent.pt < 0
    assert any(run.italic for run in reference.runs)
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "Journal of Testing" in xml
    assert "w:i" in xml


def test_docx_scanner_rejects_bare_ref_ids_and_uncited_bibliography(tmp_path) -> None:
    from docx import Document

    output = tmp_path / "leak.docx"
    doc = Document()
    doc.add_paragraph("A leaked internal reference is R001.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("Bob Jones (2025). Structured Paper Two.")
    doc.save(output)

    manifest = _manifest()
    manifest["bibliography"][1] = {**manifest["bibliography"][1], "is_cited": False}
    report = scan_docx_for_unresolved_citation_tokens(str(output), manifest)

    assert report["bare_ref_ids"] == ["R001"]
    assert report["uncited_bibliography_entries"]
    assert report["passed"] is False


def test_docx_scanner_rejects_ref_without_valid_manifest_occurrence(tmp_path) -> None:
    from docx import Document

    output = tmp_path / "unknown.docx"
    doc = Document()
    doc.add_paragraph("Claim [[cite_ref:R999]].")
    doc.save(output)

    report = scan_docx_for_unresolved_citation_tokens(str(output), _manifest())

    assert "R999" in report["unresolved_tokens"]
    assert report["passed"] is False


def test_real_docx_roundtrip_keeps_cross_section_occurrence_locators(tmp_path) -> None:
    from docx import Document

    manifest = _manifest()
    manifest["occurrences"] = [
        {"ref_id": "R001", "paper_id": "paper_1", "section_number": 1, "locator": "p. 3"},
        {"ref_id": "R001", "paper_id": "paper_1", "section_number": 2, "locator": "p. 77"},
    ]
    output = tmp_path / "roundtrip.docx"
    assert append_section_to_word_document(
        MockGenerator(), 1, "First", "Claim [[cite_ref:R001]].", str(output), citation_manifest=manifest
    )
    assert append_section_to_word_document(
        MockGenerator(), 2, "Second", "Claim [[cite_ref:R001]].", str(output), citation_manifest=manifest
    )
    paragraphs = [paragraph.text for paragraph in Document(str(output)).paragraphs]
    assert any("(Smith, 2024, p. 3)" in text for text in paragraphs)
    assert any("(Smith, 2024, p. 77)" in text for text in paragraphs)


def test_real_docx_roundtrip_disambiguates_distinct_same_author_year_papers(tmp_path) -> None:
    from docx import Document

    manifest = {
        "paper_entries": [
            {
                "paper_id": "paper_alpha",
                "paper_key": "paper_alpha",
                "title": "Alpha Study",
                "authors": ["Smith, John", "Jones, Ann", "Brown, Bob"],
                "year": "2024",
            },
            {
                "paper_id": "paper_beta",
                "paper_key": "paper_beta",
                "title": "Beta Study",
                "authors": ["Smith, John", "Taylor, Ann", "Wilson, Bob"],
                "year": "2024",
            },
        ],
        "occurrences": [
            {"ref_id": "R011", "paper_id": "paper_alpha", "section_number": 1},
            {"ref_id": "R012", "paper_id": "paper_beta", "section_number": 1},
        ],
        "bibliography": [
            {
                "entry_id": "bib_alpha",
                "paper_id": "paper_alpha",
                "paper_key": "paper_alpha",
                "citation_text": "Smith, J., Jones, A., & Brown, B. (2024). Alpha Study.",
                "is_cited": True,
            },
            {
                "entry_id": "bib_beta",
                "paper_id": "paper_beta",
                "paper_key": "paper_beta",
                "citation_text": "Smith, J., Taylor, A., & Wilson, B. (2024). Beta Study.",
                "is_cited": True,
            },
        ],
    }
    output = tmp_path / "same-author-year.docx"
    rebuild_review_docx_from_structured_artifacts(
        MockGenerator(),
        {
            "content": {
                "sections": [
                    {
                        "section_number": 1,
                        "section_title": "Synthesis",
                        "blocks": [
                            {
                                "text": (
                                    "Alpha finding [[cite_ref:R011]]; "
                                    "beta finding [[cite_ref:R012]]."
                                )
                            }
                        ],
                    }
                ]
            }
        },
        manifest,
        str(output),
    )

    paragraphs = [paragraph.text for paragraph in Document(str(output)).paragraphs]
    body = "\n".join(paragraphs)
    assert "(Smith, Jones, et al., 2024)" in body
    assert "(Smith, Taylor, et al., 2024)" in body
    assert "(Smith et al., 2024)" not in body


def test_real_docx_roundtrip_uses_full_given_names_when_initials_collide(tmp_path) -> None:
    from docx import Document

    manifest = {
        "paper_entries": [
            {
                "paper_id": "paper_john",
                "paper_key": "paper_john",
                "title": "John Study",
                "authors": ["Smith, John", "Jones, Ann", "Brown, Bob"],
                "year": "2024",
            },
            {
                "paper_id": "paper_jane",
                "paper_key": "paper_jane",
                "title": "Jane Study",
                "authors": ["Smith, Jane", "Jones, Ann", "Brown, Bob"],
                "year": "2024",
            },
        ],
        "occurrences": [
            {"ref_id": "R013", "paper_id": "paper_john", "section_number": 1},
            {"ref_id": "R014", "paper_id": "paper_jane", "section_number": 1},
        ],
        "bibliography": [
            {
                "entry_id": "bib_john",
                "paper_id": "paper_john",
                "paper_key": "paper_john",
                "citation_text": "Smith, J., Jones, A., & Brown, B. (2024). John Study.",
                "is_cited": True,
            },
            {
                "entry_id": "bib_jane",
                "paper_id": "paper_jane",
                "paper_key": "paper_jane",
                "citation_text": "Smith, J., Jones, A., & Brown, B. (2024). Jane Study.",
                "is_cited": True,
            },
        ],
    }
    output = tmp_path / "same-initial.docx"
    rebuild_review_docx_from_structured_artifacts(
        MockGenerator(),
        {
            "content": {
                "sections": [
                    {
                        "section_number": 1,
                        "section_title": "Synthesis",
                        "blocks": [
                            {
                                "text": (
                                    "John finding [[cite_ref:R013]]; "
                                    "Jane finding [[cite_ref:R014]]."
                                )
                            }
                        ],
                    }
                ]
            }
        },
        manifest,
        str(output),
    )
    body = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
    assert "(Smith, John, Jones, Ann, et al., 2024)" in body
    assert "(Smith, Jane, Jones, Ann, et al., 2024)" in body


def test_real_docx_roundtrip_keeps_year_suffix_in_body_and_references(tmp_path) -> None:
    from docx import Document

    manifest = {
        "paper_entries": [
            {
                "paper_id": "paper_a",
                "paper_key": "paper_a",
                "title": "Alpha Study",
                "authors": ["Smith, John"],
                "year": "2024",
                "journal": "Journal A",
            },
            {
                "paper_id": "paper_b",
                "paper_key": "paper_b",
                "title": "Beta Study",
                "authors": ["Smith, John"],
                "year": "2024",
                "journal": "Journal B",
            },
        ],
        "occurrences": [
            {"ref_id": "R015", "paper_id": "paper_a", "section_number": 1},
            {"ref_id": "R016", "paper_id": "paper_b", "section_number": 1},
        ],
        "bibliography": [
            {
                "entry_id": "bib_a",
                "paper_id": "paper_a",
                "paper_key": "paper_a",
                "citation_text": "Smith, J. (2024). Alpha Study. Journal A.",
                "is_cited": True,
            },
            {
                "entry_id": "bib_b",
                "paper_id": "paper_b",
                "paper_key": "paper_b",
                "citation_text": "Smith, J. (2024). Beta Study. Journal B.",
                "is_cited": True,
            },
        ],
    }
    output = tmp_path / "same-author-suffix.docx"
    rebuild_review_docx_from_structured_artifacts(
        MockGenerator(),
        {
            "content": {
                "sections": [
                    {
                        "section_number": 1,
                        "section_title": "Synthesis",
                        "blocks": [
                            {
                                "text": (
                                    "Alpha [[cite_ref:R015]]; "
                                    "beta [[cite_ref:R016]]."
                                )
                            }
                        ],
                    }
                ]
            }
        },
        manifest,
        str(output),
    )
    paragraphs = [paragraph.text for paragraph in Document(str(output)).paragraphs]
    body = "\n".join(paragraphs)
    assert "(Smith, 2024a)" in body
    assert "(Smith, 2024b)" in body
    assert any("(2024a)" in paragraph for paragraph in paragraphs)
    assert any("(2024b)" in paragraph for paragraph in paragraphs)
